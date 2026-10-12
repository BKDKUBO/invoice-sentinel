"""Build InvoiceSentinel proposal: PDF (Chrome print of proposal.html) and DOCX (editable copy).
Usage: python build.py  (needs playwright + python-docx + bs4; Chrome at /usr/bin/google-chrome)"""
import asyncio, pathlib, re, sys
from bs4 import BeautifulSoup, NavigableString, Tag
HERE = pathlib.Path(__file__).resolve().parent
OUT_PDF = HERE.parent / "InvoiceSentinel-Proposal.pdf"
OUT_DOCX = HERE / "InvoiceSentinel-Proposal.docx"
FOOT = "InvoiceSentinel · Google Cloud AI Builder Cup 2026 · BFSI: Intelligent Risk, Fraud &amp; Financial Experiences"

def build_html():
    soup = BeautifulSoup((HERE / "proposal.html").read_text(), "html.parser")
    pages = soup.select("section.page")
    for i, s in enumerate(pages, 1):
        if "cover" in s.get("class", []):
            continue
        f = soup.new_tag("div", attrs={"class": "foot nodoc"})
        f.append(BeautifulSoup(f"<span>{FOOT}</span><span>{i} / {len(pages)}</span>", "html.parser"))
        s.append(f)
    out = HERE / "_print.html"
    out.write_text(str(soup))
    return out

async def pdf(html):
    from playwright.async_api import async_playwright
    async with async_playwright() as p:
        b = await p.chromium.launch(executable_path="/usr/bin/google-chrome")
        pg = await b.new_page(viewport={"width": 1280, "height": 720})
        await pg.goto(html.as_uri(), wait_until="networkidle")
        await pg.evaluate("document.fonts.ready")
        # overflow check: content bottom must stay above footer
        over = await pg.evaluate("""() => [...document.querySelectorAll('section.page')].map((s,i)=>{
            const r=s.getBoundingClientRect(); let m=0;
            s.querySelectorAll('*').forEach(e=>{ if(e.closest('.foot')||e.closest('.band')) return; const b=e.getBoundingClientRect(); if(b.height>0) m=Math.max(m,b.bottom-r.top);});
            return [i+1, Math.round(m)];})""")
        print("content bottom per page (limit ~690):", over)
        await pg.pdf(path=str(OUT_PDF), prefer_css_page_size=True, print_background=True)
        await b.close()

# ---------- DOCX ----------
def build_docx():
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.section import WD_ORIENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    soup = BeautifulSoup((HERE / "proposal.html").read_text(), "html.parser")
    d = Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("left_margin", "right_margin"): setattr(sec, m, Inches(0.8))
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
    NAVY = RGBColor(0x0B, 0x1F, 0x3A)
    for name, size in (("Title", 28), ("Heading 1", 18), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        s = d.styles[name]; s.font.name = "Arial"; s.font.size = Pt(size); s.font.color.rgb = NAVY; s.font.bold = True

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

    def runs(par, node, bold=False, ital=False, mono=False, hl=False):
        for ch in node.children:
            if isinstance(ch, NavigableString):
                t = re.sub(r"\s+", " ", str(ch))
                if not t: continue
                r = par.add_run(t); r.bold = bold or None; r.italic = ital or None
                if mono: r.font.name = "Courier New"; r.font.size = Pt(9.5)
                if hl:
                    from docx.enum.text import WD_COLOR_INDEX
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif isinstance(ch, Tag):
                if ch.name == "br": par.add_run().add_break(); continue
                if ch.name == "a" and ch.get("href", "").startswith("http"):
                    r = par.add_run(ch["href"]); r.font.color.rgb = RGBColor(0x1A, 0x56, 0xB8); r.bold = bold or None; continue
                cls = ch.get("class", [])
                runs(par, ch, bold or ch.name in ("b", "strong") or "num" in cls, ital or ch.name in ("i", "em"),
                     mono or ch.name == "code" or "mono" in cls, hl or "fill" in cls or "fillc" in cls)

    def para(node, style=None, prefix=None):
        p = d.add_paragraph(style=style)
        if prefix: p.add_run(prefix).bold = True
        runs(p, node)
        p.paragraph_format.space_after = Pt(4)
        return p

    def table(node):
        rows = node.find_all("tr")
        ncol = max(len(r.find_all(["td", "th"])) for r in rows)
        t = d.add_table(rows=0, cols=ncol); t.style = "Table Grid"
        for r in rows:
            cells = r.find_all(["td", "th"]); row = t.add_row().cells
            for j, c in enumerate(cells):
                p = row[j].paragraphs[0]; runs(p, c, bold=c.name == "th")
                for rr in p.runs: rr.font.size = Pt(9.5)
                if c.name == "th":
                    shade(row[j], "0B1F3A")
                    for rr in p.runs: rr.font.color.rgb = RGBColor(255, 255, 255)
        d.add_paragraph()

    INLINE = {"a", "b", "strong", "i", "em", "span", "code", "br", "small"}
    def walk(node):
        pending = []
        def flush():
            if pending and any((isinstance(x, NavigableString) and x.strip()) or isinstance(x, Tag) for x in pending):
                holder = BeautifulSoup("<p></p>", "html.parser").p
                for x in pending: holder.append(x.__copy__() if isinstance(x, Tag) else NavigableString(str(x)))
                if holder.get_text(strip=True): para(holder)
            pending.clear()
        for ch in list(node.children):
            if isinstance(ch, NavigableString):
                if node.name == "div" and str(ch).strip(): pending.append(ch)
                continue
            if node.name == "div" and ch.name in INLINE and not set(ch.get("class", [])) & {"nodoc", "pill", "fill"} - {"fill"}:
                if "sev" in ch.get("class", []) or "tag" in ch.get("class", []) or ch.name != "span" or True:
                    pending.append(ch); continue
            flush()
            cls = ch.get("class", [])
            if "nodoc" in cls or "foot" in cls: continue
            n = ch.name
            if n == "h1": d.add_heading(ch.get_text(" ", strip=True), level=1)
            elif n == "h2": d.add_heading(ch.get_text(" ", strip=True), level=2)
            elif n == "h3": d.add_heading(ch.get_text(" ", strip=True), level=3)
            elif n == "p":
                if "kicker" in cls:
                    p = d.add_paragraph(); r = p.add_run(ch.get_text(strip=True).upper()); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x1A, 0x56, 0xB8); r.bold = True
                else: para(ch)
            elif n in ("ul", "ol"):
                for li in ch.find_all("li", recursive=False):
                    para(li, style="List Number" if n == "ol" else "List Bullet")
            elif n == "table": table(ch)
            elif n == "figure":
                img = ch.find("img")
                if img:
                    d.add_picture(str(HERE / img["src"]), width=Inches(6.9 if "architecture" in img["src"] else 6.4))
                cap = ch.find("figcaption")
                if cap:
                    p = para(cap); [setattr(r.font, "size", Pt(9)) for r in p.runs]
            elif "callout" in cls:
                p = para(ch); p.paragraph_format.left_indent = Inches(0.25)
                for r in p.runs: r.font.color.rgb = RGBColor(0x6B, 0x4E, 0x00) if "blue" not in cls and "green" not in cls else RGBColor(0x0B, 0x1F, 0x3A)
            elif "stat" in cls:
                p = d.add_paragraph(style="List Bullet")
                for part in ch.find_all("div", recursive=False):
                    runs(p, part, bold="num" in part.get("class", [])); p.add_run("  ")
            elif "step" in cls:
                pp = d.add_paragraph(style="List Bullet")
                n_ = ch.find(class_="n"); b_ = ch.find("b")
                pp.add_run((n_.get_text(strip=True) if n_ else "") + ": ").font.size = Pt(9)
                pp.add_run((b_.get_text(strip=True) if b_ else "") + ". ").bold = True
                pp.add_run("".join(str(x) for x in ch.children if isinstance(x, NavigableString)).strip())
            elif "crit" in cls:
                p = d.add_paragraph(); r = p.add_run("Rubric: " + " · ".join(x.get_text(strip=True) for x in ch.find_all("span"))); r.italic = True; r.font.size = Pt(9)
            elif "row" in cls:
                sp = ch.find_all(["span", "b"])
                pp = d.add_paragraph(style="List Bullet"); pp.add_run(sp[0].get_text(strip=True) + ": "); pp.add_run(sp[1].get_text(strip=True)).bold = True
            elif "k" in cls:
                d.add_heading(ch.get_text(strip=True), level=3)
            else: walk(ch)
        flush()

    for i, s in enumerate(soup.select("section.page")):
        if i > 0: d.add_page_break()
        walk(s)
    d.save(OUT_DOCX); print("docx:", OUT_DOCX)

if __name__ == "__main__":
    asyncio.run(pdf(build_html())); print("pdf:", OUT_PDF)
    build_docx()
