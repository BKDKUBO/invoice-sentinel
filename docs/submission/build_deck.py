"""Builds InvoiceSentinel-BKD-KUBO-Deck.pptx on the AI Builder Cup template (backgrounds rendered from the official template PDF)."""
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
import pathlib
H=pathlib.Path(__file__).resolve().parent; I=H/"img"
URL="https://invoicesentinel.web.app"; RUNURL="https://invoicesentinel-web-1034883854575.asia-southeast1.run.app"
NAVY=RGBColor(0x0B,0x1C,0x2C); GOLD=RGBColor(0xA2,0x7A,0x06); INK=RGBColor(0x20,0x21,0x24); MUTE=RGBColor(0x5F,0x63,0x68); LINE=RGBColor(0xDA,0xDC,0xE0)
SEV={"critical":RGBColor(0xCE,0x11,0x26),"high":RGBColor(0xE8,0x76,0x1C),"low":RGBColor(0x2E,0x9D,0x5B)}
FONT="Arial"
P=Presentation(); P.slide_width=Pt(720); P.slide_height=Pt(405); BL=P.slide_layouts[6]
def slide(bg):
    s=P.slides.add_slide(BL); s.shapes.add_picture(str(I/bg),0,0,P.slide_width,P.slide_height); return s
def tb(s,x,y,w,h,paras,anchor=MSO_ANCHOR.TOP):
    t=s.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h)); f=t.text_frame; f.word_wrap=True; f.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(f,m,Pt(2))
    for i,p in enumerate(paras):
        if isinstance(p,str): p=(p,)
        txt,size,bold,color,align=(list(p)+[None]*5)[:5]
        para=f.paragraphs[0] if i==0 else f.add_paragraph(); para.space_after=Pt(3)
        if align: para.alignment=align
        parts=txt.split("**")
        for k,part in enumerate(parts):
            if not part: continue
            r=para.add_run(); r.text=part; r.font.name=FONT; r.font.size=Pt(size or 10.5); r.font.bold=bool(bold) or k%2==1; r.font.color.rgb=color or INK
    return t
def title(s,text): tb(s,25.6,52,670,26,[(text,14,True,INK)])
def bul(items,size=10.5): return [("•  "+x,size) for x in items]
def img(s,name,x,y,w=None,h=None,border=True):
    p=s.shapes.add_picture(str(I/name),Pt(x),Pt(y),Pt(w) if w else None,Pt(h) if h else None)
    if border: p.line.color.rgb=LINE; p.line.width=Pt(0.75)
    return p
def card(s,x,y,w,h,head,body,fill=RGBColor(0xF8,0xF9,0xFA),hc=NAVY,size=9.5):
    r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h)); r.adjustments[0]=0.08
    r.fill.solid(); r.fill.fore_color.rgb=fill; r.line.color.rgb=LINE; r.line.width=Pt(0.75); r.shadow.inherit=False
    tb(s,x+6,y+5,w-12,h-10,[(head,11,True,hc)]+[(b,size) for b in body])
def table(s,x,y,w,rows,colw,size=8.5,rowh=13):
    t=s.shapes.add_table(len(rows),len(rows[0]),Pt(x),Pt(y),Pt(w),Pt(rowh*len(rows))).table
    for j,cw in enumerate(colw): t.columns[j].width=Pt(cw)
    for i,row in enumerate(rows):
        t.rows[i].height=Pt(rowh)
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=""; c.margin_left=c.margin_right=Pt(4); c.margin_top=c.margin_bottom=Pt(1)
            c.fill.solid(); c.fill.fore_color.rgb=NAVY if i==0 else (RGBColor(0xF8,0xF9,0xFA) if i%2 else RGBColor(0xFF,0xFF,0xFF))
            r=c.text_frame.paragraphs[0].add_run(); r.text=str(v); r.font.name=FONT; r.font.size=Pt(size); r.font.bold=(i==0)
            r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF) if i==0 else INK
    return t

# 1 Team Details
s=slide("bg_team.jpg")
tb(s,17,226,300,30,[("Team Details",16,True,INK)])
tb(s,40,268,470,110,[("a.   Team name: **BKD KUBO**",12),("b.   Team leader name: **Manuel III Ocaña**",12),("c.   Problem Statement: **BFSI: Intelligent Risk, Fraud & Financial Experiences**",12),
   ("Members: Manuel III Ocaña — Product Lead & Solution Architect · Elson Mojica — Process & QA Lead (BFSI Operations)",9,False,MUTE),("Solution: InvoiceSentinel, a pre-payment AP fraud & risk desk for PH SMEs",9,False,MUTE)])
img(s,"teamlogo.jpg",600,238,92,None,border=False)

# 3 Brief about the idea
s=slide("bg_content.jpg"); title(s,"Brief about the idea")
tb(s,28,92,340,300,[("**InvoiceSentinel** checks every supplier invoice before it is paid and tells the AP analyst **PAY, HOLD or VERIFY**, with evidence.",11),("",4)]+bul([
 "**Problem:** PH ghost receipts: the BIR alleges ₱25.5B in tax losses from 4 ghost firms (2019–21) and ₱1.41B against 23 buyer firms (2025). 68% of organizations saw an AP fraud attempt in 2024 (Basware).",
 "**Who:** SMEs with no fraud team, where one clerk reviews PDFs by eye.",
 "**How:** Gemini 3.8 Flash reads the invoice; deterministic code checks vendor master, PO, TIN, BIR RR 7-2024 and split invoices; a Google ADK agent explains in the user's language.",
 "**Human in the loop:** Approve / Escalate / Request docs, with an audit trail. Nothing is paid or sent automatically."],10))
img(s,"dashboard.jpg",380,96,312)
tb(s,380,296,312,14,[("Live dashboard on Cloud Run: KPI cards, risk trend, flag breakdown",8,False,MUTE)])

# 4 Opportunities
s=slide("bg_content.jpg"); title(s,"Opportunities")
card(s,28,94,212,290,"How is it different from existing ideas?",["•  OCR + rules tools need templates and can't judge \"new bank, don't call us\".","•  Generic AI chat over a PDF has no vendor master, no PO, no audit trail.","•  Opaque fraud scores can't be defended to auditors.","•  InvoiceSentinel combines template-free Gemini reading with deterministic checks, evidence citations and a human decision."])
card(s,254,94,212,290,"How will it solve the problem?",["•  Flags the patterns that drain SMEs: ghost vendors, bank-account changes (BEC), duplicates, overbilling, PO mismatches, split invoices.","•  Ranks the inbox by a deterministic 0–100 score, so high-risk invoices are reviewed first.","•  Drafts a verification email to the contact on file, never the one on the invoice.","•  BIR RR 7-2024 checklist keeps ghost receipts out of the books."])
card(s,480,94,212,290,"USP of the proposed solution",["•  **Code for facts, AI for judgment:** same input, same score (14 samples × 3 runs, zero variance).","•  **Explainable agent:** Google ADK agent streams each tool step live and cites evidence IDs.","•  **PH-first:** TIN, 12% VAT, BIR RR 7-2024, GCash/personal payees, Tagalog answers and emails.","•  **Affordable:** ~₱0.31 Gemini cost per invoice."],fill=RGBColor(0xFE,0xF7,0xE0),hc=GOLD)

# 5 Features
s=slide("bg_content.jpg"); title(s,"List of features offered by the solution")
F=[("Multimodal extraction","Gemini 3.8 Flash reads photos/PDFs into structured JSON (T=0, schema)."),("Vendor master & PO matching","Vendor, TIN, bank account, PO, quantities and prices; bank change = critical."),
("TIN validation","Format and branch-code checks, compared with the master record."),("BIR RR 7-2024 checklist","Required invoice elements with section references; advisory only."),
("Split-invoice detection","Same vendor, close dates, combined total over the approval limit."),("Deterministic risk score","Severity weights + floors; ranked inbox and ₱-at-risk KPIs."),
("Ask the case (Google ADK)","7 read-only tools; live streamed steps; answers end with PAY / HOLD / VERIFY."),("Multilingual + verification email","Replies in English, Tagalog, etc.; email to the contact on file, never auto-sent."),
("Try a sample","14 fictional PH invoices, one click, no files needed."),("Decisions & audit trail","Approve / Escalate / Request docs with notes; timestamped log."),
("Vendors & POs page","Edit the master data the checks use; approval limit in Settings."),("Accessible, responsive UI","AA contrast, keyboard & focus states, ARIA labels, mobile layouts.")]
for k,(h,b) in enumerate(F):
    x=28+(k%3)*224; y=92+(k//3)*75
    card(s,x,y,214,68,h,[b],size=9)

# 6 Process flow
s=slide("bg_content.jpg"); title(s,"Process flow diagram or Use-case diagram")
steps=[("1  Upload","Photo, scan or PDF, or one click on Try a sample"),("2  Gemini reads","Fields, line items, payee & bank details as JSON"),("3  Code checks","Vendor master, PO, TIN, BIR RR 7-2024, split, VAT math → score"),("4  Agent explains","Ask in any language; ADK agent cites evidence, drafts email"),("5  Human decides","Approve / Escalate / Request docs → audit trail")]
for k,(h,b) in enumerate(steps):
    x=28+k*136; r=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Pt(x),Pt(100),Pt(118),Pt(92)); r.adjustments[0]=0.1
    r.fill.solid(); r.fill.fore_color.rgb=NAVY if k in (2,3) else RGBColor(0xF8,0xF9,0xFA); r.line.color.rgb=GOLD if k in (2,3) else LINE; r.shadow.inherit=False
    w=RGBColor(0xFF,0xFF,0xFF) if k in (2,3) else INK
    tb(s,x+5,104,108,84,[(h,11,True,RGBColor(0xFC,0xD1,0x16) if k in (2,3) else NAVY),(b,9,False,w)])
    if k<4:
        a=s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Pt(x+120),Pt(140),Pt(14),Pt(12)); a.fill.solid(); a.fill.fore_color.rgb=GOLD; a.line.fill.background()
tb(s,28,206,664,16,[("Use case: AP analyst at an SME receives invoice from known vendor “Northstar IT” (sample 09)",11,True,NAVY)])
tb(s,28,226,664,160,bul(["Upload → Gemini extracts vendor, TIN, PO, totals and the bank account printed on the invoice.",
 "Vendor master check: everything matches except the bank account (invoice Metrobank vs master BPI) → UNUSUAL_PAYMENT_CHANGE, critical, score 90.",
 "Analyst asks in Tagalog “Ligtas ba itong bayaran?” → agent calls vendor, PO, history, TIN and BIR tools → answers HOLD with evidence.",
 "Analyst drafts a Tagalog verification email to the master contact, calls the supplier, and escalates; the decision is logged."],10))

# 7 Wireframes
s=slide("bg_content.jpg"); title(s,"Wireframes/Mock diagrams of the proposed solution (optional)")
img(s,"upload.jpg",28,92,300); tb(s,28,288,300,14,[("Upload: guided stepper + Try a sample grid",8,False,MUTE)])
img(s,"case.jpg",340,92,262); tb(s,340,282,262,14,[("Case: fields, risk cards, matching, decision panel",8,False,MUTE)])
img(s,"mobile.jpg",618,92,None,180); tb(s,614,274,90,14,[("Mobile layout",8,False,MUTE)])
tb(s,28,306,560,80,bul(["Implemented in the live prototype (screens shown are real, not mockups).","Flow: Dashboard → Check an invoice → Case → Decide; Vendors & POs page for master data."],9.5))

# 8 Architecture
s=slide("bg_content.jpg"); title(s,"Architecture diagram of the proposed solution")
img(s,"architecture.png",28,90,664,None,border=False)
fb=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Pt(30),Pt(98),Pt(104),Pt(40)); fb.fill.solid(); fb.fill.fore_color.rgb=RGBColor(0xFE,0xF7,0xE0); fb.line.color.rgb=RGBColor(0xF5,0x82,0x0B); fb.line.width=Pt(1.25); fb.shadow.inherit=False
tb(s,32,100,100,36,[("Firebase Hosting",9,True,NAVY,PP_ALIGN.CENTER),("invoicesentinel.web.app → rewrites to Cloud Run web",7,False,MUTE,PP_ALIGN.CENTER)])
tb(s,28,322,664,72,bul(["Firebase Hosting serves invoicesentinel.web.app and rewrites all paths to the Cloud Run web service; the browser calls the API directly (no 60 s Hosting limit on agent streams).",
 "Stateless Next.js and FastAPI containers on Cloud Run (asia-southeast1); state in Firestore and a private Cloud Storage bucket.",
 "Gemini key in Secret Manager; one-command deploy.sh builds both images with Cloud Build."],9.5))

# 9 Technologies
s=slide("bg_content.jpg"); title(s,"Technologies to be used in the solution")
T=[("Google AI","Gemini 3.8 Flash (google-genai SDK): multimodal read, JSON schema, T=0 · Google ADK 2.11: LlmAgent with 7 tools, streamed events"),
("Google Cloud & Firebase","Firebase Hosting (invoicesentinel.web.app → Cloud Run) · Cloud Run (web + API, min-instances 1) · Firestore · Cloud Storage · Secret Manager · Cloud Build · Artifact Registry"),
("Backend","Python 3.12 · FastAPI · Pydantic · deterministic rule engine (vendor/PO matching, TIN, BIR RR 7-2024, split detection, scoring) · NDJSON streaming"),
("Frontend","Next.js 14 · React · TypeScript · Tailwind CSS · Inter · custom SVG charts · responsive, WCAG AA-minded UI"),
("Quality","Stability harness (14 fictional samples × 3 runs) · sample generator · Playwright screenshots")]
for k,(h,b) in enumerate(T):
    y=94+k*58; card(s,28,y,664,50,h,[b],size=10)

# 10 Cost
s=slide("bg_content.jpg"); title(s,"Estimated implementation cost (optional)")
card(s,28,94,200,90,"~₱0.31 per invoice",["Gemini 3.8 Flash at 2026 promo pricing (~₱0.62 from Jan 2027). FX ₱62.77/US$."],fill=RGBColor(0xFE,0xF7,0xE0),hc=GOLD)
table(s,244,96,448,[["Invoices / month","Typical user","Model cost 2026","From 2027"],["200","Small SME","~₱62","~₱124"],["2,000","Mid-market finance team","~₱623","~₱1,246"],["50,000","Shared services / bank SME portfolio","~₱16,000","~₱31,000"]],[90,190,84,84],size=9.5,rowh=20)
tb(s,28,200,664,190,bul(["**Infrastructure:** Cloud Run scales to zero; at SME volumes usage fits the free tier. For judging, one warm instance per service (min-instances 1): ~US$23 (~₱1,450) per month; set to 0 afterwards.",
 "**Storage:** Firestore and Cloud Storage cost almost nothing at prototype scale.",
 "**Scale levers:** Gemini Batch API for overnight payment runs at half the token price; Flash by default.",
 "Source: Gemini Developer API pricing (updated Oct 9, 2026) and Cloud Run pricing. Per-invoice token use is an estimate."],10))

# 11 Snapshots
s=slide("bg_content.jpg"); title(s,"Snapshots of the prototype")
SN=[("s_vendor.jpg","Vendor master: bank account mismatch (sample 09)"),("s_agent.jpg","Google ADK agent steps streaming live"),("s_answer.jpg","Tagalog answer: HOLD with cited evidence"),
("s_email.jpg","Tagalog verification email to the contact on file"),("s_bir.jpg","BIR RR 7-2024 checklist + TIN validation"),("s_split.jpg","Sample 13 flagged as a split of sample 12")]
for k,(f,c) in enumerate(SN):
    x=28+(k%3)*224; y=90+(k//3)*152
    img(s,f,x,y,214,122); tb(s,x,y+124,214,14,[(c,8,False,MUTE)])

# 12 Performance
s=slide("bg_content.jpg"); title(s,"Prototype Performance report/Benchmarking")
EV=[("01",0,"low"),("02",0,"low"),("03",100,"critical"),("04",70,"high"),("05",90,"critical"),("06",75,"high"),("07",70,"high"),("08",20,"low"),("09",90,"critical"),("10",70,"high"),("11",70,"high"),("12",0,"low"),("13",70,"high"),("14",80,"high")]
cd=CategoryChartData(); cd.categories=[e[0] for e in EV]; cd.add_series("Risk score (3 runs identical)",[max(e[1],1) for e in EV])
ch=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Pt(24),Pt(88),Pt(330),Pt(170),cd).chart
ch.has_legend=False; ch.has_title=True; ch.chart_title.text_frame.text="Risk score per sample (14 samples × 3 runs, zero spread)"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size=Pt(9); ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold=True
ch.value_axis.maximum_scale=100; ch.value_axis.tick_labels.font.size=Pt(8); ch.category_axis.tick_labels.font.size=Pt(8); ch.value_axis.has_major_gridlines=True
ch.value_axis.major_gridlines.format.line.color.rgb=LINE; ch.plots[0].gap_width=60
for k,e in enumerate(EV):
    pt=ch.plots[0].series[0].points[k]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb=SEV[e[2]]
table(s,24,266,330,[["Metric (live, Cloud Run)","Result"],["Score stability, 14 samples × 3 runs","0 spread, 14/14 in band"],["Upload → verdict (Gemini + checks)","~18–24 s"],["Agent: first streamed step","~5 s"],["Agent: full grounded answer","~18–37 s"],["Verification email draft","~20–25 s"]],[200,130],size=8.5,rowh=19)
table(s,366,92,326,[["Flag (expected samples)","Detected"],["GHOST_VENDOR (03, 14)","2 / 2"],["UNUSUAL_PAYMENT_CHANGE (05, 09)","2 / 2"],["TIN_INVALID (03, 14)","2 / 2"],["MATH_ERROR (04)","1 / 1"],["OVERBILLING (06)","1 / 1"],["DUPLICATE_INVOICE (07)","1 / 1"],["PO_MISMATCH (06, 10)","2 / 2"],["MISSING_PO (11)","1 / 1"],["SPLIT_INVOICE (13 after 12)","1 / 1"],["Clean invoices without fraud flags (01, 02, 08, 12)","4 / 4"]],[250,76],size=8.5,rowh=19.5)
tb(s,366,312,326,80,[("Fictional PH test set, gemini-3.8-flash, Oct 12–13, 2026. Sample 08 (blurry photo) gets only a low-severity TIN-mismatch flag, not a fraud flag. Small test set: a labeled real-world evaluation set is on the roadmap.",8,False,MUTE)])

# 13 Additional
s=slide("bg_content.jpg"); title(s,"Additional Details/Future Development (if any)")
R=[("Q1 2027 · Pilot-ready",["Sign-in with analyst / approver roles (Identity Platform)","Email-inbox intake of supplier bills","CSV import of vendor master and POs","Labeled PH evaluation set; precision/recall per flag"]),
("Q2 2027 · Integrations",["QuickBooks / Xero bill sync; hold payment on Escalate","Batch review of payment runs (Gemini Batch API)","Bank-account change history per vendor"]),
("Later · Ecosystem",["Pre-transfer check API for bank SME portals","Claims-document review for insurers","Tagalog / Cebuano UI","BIR ATP / series lookups as signals"])]
for k,(h,xs) in enumerate(R): card(s,28+k*224,94,214,200,h,["•  "+x for x in xs],size=9.5)
tb(s,28,306,664,80,[("**Guardrails stay fixed:** a person decides every payment, every flag carries evidence, nothing is auto-sent, and the product never claims to certify BIR compliance. Paid-tier Gemini API: content is not used to train Google's models.",10)])

# 14 Links
s=slide("bg_content.jpg"); title(s,"Provide links to your:")
tb(s,46,94,640,200,[("•  GitHub Public Repository",12,True),("     https://github.com/BKDKUBO/invoice-sentinel",11,False,RGBColor(0x1A,0x56,0xB8)),("",6),
 ("•  Demo Video Link (3 Minutes)",12,True),("     https://drive.google.com/file/d/1mIB4FfR1lKYbUxSovZpBJaOMQzojOQUn/view",11,False,GOLD),("",6),
 ("•  Final Product Link",12,True),("     "+URL,11,False,RGBColor(0x1A,0x56,0xB8)),("     (Firebase Hosting → Cloud Run; direct: "+RUNURL+")",8.5,False,MUTE),("",6),
 ("Try it: Check an invoice → Try a sample → 09 (known vendor, new bank). Reset demo restores seed data. All data is fictional.",9,False,MUTE)])

# 16 Thank you (template closing slide)
slide("bg_thanks.jpg")
out=H/"InvoiceSentinel-BKD-KUBO-Deck.pptx"; P.save(out); print(out)
